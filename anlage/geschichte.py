"""Geschichtsdaten für den Rückblick: einmalig holen, fortsetzbar.

- daten/geschichte/preise.json            Monatspreise ab 1960 (Weltbank Pink Sheet; Rückfall FRED ab 1992)
- daten/geschichte/ernte.json             Palmöl-Produktion je Land und Wirtschaftsjahr (USDA PSD)
- daten/geschichte/regen/<punkt>.json.gz  Tagesregen (Ostpunkte ganzjährig, Westpunkte nur Oktober bis Dezember)
- daten/geschichte/wind.json              mittlerer Ost-West-Wind im Mai und Juni je Jahr
- daten/geschichte/status.json            was geholt ist und was fehlt

Open-Meteo rechnet lange Zeiträume mehrfach an. Der Regen braucht deshalb etwa
12 000 gewichtete Abrufe; bei einem Tageslimit von 10 000 sind das zwei Läufe an
zwei Tagen. Jeder Lauf stoppt sauber vor seinem Budget und macht beim nächsten
Start dort weiter, wo er aufgehört hat.

Aufruf: python -m anlage.geschichte [--nur preise,ernte,regen,wind] [--nur-offene] [--budget-gewicht 9000]
        [--budget-minuten 150]
Mit --nur-offene holt der Lauf nur, was noch nicht fertig ist, und tut nichts, wenn alles da ist
(für den täglichen Zeitplan, bis die Geschichte vollständig ist).
"""
from __future__ import annotations

import argparse
import calendar
import csv
import gzip
import io
import json
import os
import re
import time
import zipfile
from datetime import date, timedelta

from . import konfig, netz, speicher
from .klimatologie import Abbruch
from .netz import AbrufFehler
from .quellen import kurse, openmeteo


def _norm(text) -> str:
    return " ".join(str(text).split()).lower() if text is not None else ""


def _slug(name: str) -> str:
    ersatz = {"ä": "ae", "ö": "oe", "ü": "ue", "ß": "ss"}
    text = "".join(ersatz.get(z, z) for z in name.lower())
    return re.sub(r"[^a-z0-9]+", "-", text).strip("-")


def _datum(wert) -> date:
    return wert if isinstance(wert, date) else date.fromisoformat(str(wert))


class Budget:
    """Zählt gewichtete Open-Meteo-Abrufe und Zeit; hält an, bevor eine Grenze überschritten wird."""

    def __init__(self, gewicht: float, minuten: float, pro_minute: float):
        self.grenze, self.verbraucht = gewicht, 0.0
        self.ende = time.monotonic() + minuten * 60
        self.pro_minute = pro_minute

    def pruefe(self, gewicht: float, was: str) -> None:
        if self.verbraucht + gewicht > self.grenze:
            raise Abbruch(f"Abrufbudget für heute erreicht vor „{was}“. Morgen einfach noch einmal starten.")
        dauer = 60.0 * gewicht / self.pro_minute if self.pro_minute else 0
        if time.monotonic() + dauer > self.ende:
            raise Abbruch(f"Zeitbudget erreicht vor „{was}“. Einfach noch einmal starten.")

    def verbuche(self, gewicht: float) -> None:
        self.verbraucht += gewicht
        openmeteo.drossel(gewicht, self.pro_minute)


# ------------------------------------------------------------------ Preise
def lies_weltbank(inhalt: bytes, spalten: dict[str, str]) -> tuple[dict[str, list[dict]], list[str]]:
    """Liest das Blatt mit den Monatspreisen aus dem Pink Sheet.

    spalten: {Kennung: Spaltenname im Pink Sheet}. Gibt die Reihen und die Liste der
    nicht gefundenen Spalten zurück. Datumsspalte im Format 1960M01.
    """
    import openpyxl
    mappe = openpyxl.load_workbook(io.BytesIO(inhalt), read_only=True, data_only=True)
    blaetter = [n for n in mappe.sheetnames if "month" in n.lower()] or mappe.sheetnames
    for name in blaetter:
        zeilen = [list(z) for z in mappe[name].iter_rows(values_only=True)]
        kopf_nr, treffer = None, {}
        for nr, zeile in enumerate(zeilen[:30]):
            normiert = [_norm(z) for z in zeile]
            gefunden = {}
            for kennung, spalte in spalten.items():
                ziel = _norm(spalte)
                if ziel in normiert:
                    gefunden[kennung] = normiert.index(ziel)
                else:
                    passend = [i for i, z in enumerate(normiert) if z.startswith(ziel)]
                    if passend:
                        gefunden[kennung] = passend[0]
            if len(gefunden) > len(treffer):
                kopf_nr, treffer = nr, gefunden
        if not treffer:
            continue
        reihen = {k: [] for k in treffer}
        for zeile in zeilen[kopf_nr + 1:]:
            m = re.fullmatch(r"(\d{4})M(\d{2})", str(zeile[0]).strip()) if zeile and zeile[0] else None
            if not m:
                continue
            datum = f"{m.group(1)}-{m.group(2)}-01"
            for kennung, spalte in treffer.items():
                wert = zeile[spalte] if spalte < len(zeile) else None
                if isinstance(wert, (int, float)) and not isinstance(wert, bool):
                    reihen[kennung].append({"datum": datum, "wert": round(float(wert), 4)})
        fehlend = [k for k in spalten if k not in treffer or not reihen.get(k)]
        return {k: v for k, v in reihen.items() if v}, fehlend
    return {}, list(spalten)


def _weltbank_urls(g: dict) -> list[str]:
    urls = list(g["weltbank"].get("urls") or [])
    try:
        seite = netz.hole_text(g["weltbank"]["seite"])
        for treffer in re.findall(r'href="([^"]*CMO-Historical-Data-Monthly\.xlsx)"', seite):
            url = treffer if treffer.startswith("http") else "https://www.worldbank.org" + treffer
            if url not in urls:
                urls.append(url)
    except AbrufFehler:
        pass
    return urls


def preise(k: dict, melde=print) -> dict:
    g = k["quellen"]["geschichte"]
    spalten = {kennung: e["weltbank"] for kennung, e in g["preise"].items()}
    reihen, quelle, fehler = {}, None, []
    for url in _weltbank_urls(g):
        try:
            reihen, fehlend = lies_weltbank(netz.hole_bytes(url), spalten)
        except Exception as f:
            fehler.append(f"{url}: {type(f).__name__}: {f}")
            continue
        if reihen:
            quelle = f"Weltbank Pink Sheet ({url.rsplit('/', 1)[-1]})"
            if fehlend:
                fehler.append("im Pink Sheet nicht gefunden: " + ", ".join(fehlend))
            break
    herkunft = {kennung: quelle for kennung in reihen}
    for kennung, eintrag in g["preise"].items():
        if kennung in reihen or not eintrag.get("fred"):
            continue
        try:
            reihen[kennung] = kurse.hole_fred(k["quellen"]["fred"]["url"], eintrag["fred"])
            herkunft[kennung] = f"FRED {eintrag['fred']} (erst ab 1992)"
        except Exception as f:
            fehler.append(f"FRED {eintrag['fred']}: {f}")
    # Eine Reihe, die später beginnt als die schon gespeicherte, ersetzt sie nicht.
    alt = speicher.lies_json(konfig.GESCHICHTE / "preise.json", {}) or {}
    for kennung, reihe_alt in (alt.get("reihen") or {}).items():
        neu = reihen.get(kennung)
        if reihe_alt and (not neu or neu[0]["datum"] > reihe_alt[0]["datum"]):
            reihen[kennung] = reihe_alt
            herkunft[kennung] = f"{(alt.get('herkunft') or {}).get(kennung, 'früherer Abruf')}, Stand {alt.get('stand')}"
            fehler.append(f"{kennung}: heute keine gleich lange Reihe, der Abruf vom {alt.get('stand')} bleibt")
    if not reihen:
        raise AbrufFehler("Keine Preisreihe lesbar. " + " | ".join(fehler))
    daten = {"stand": date.today().isoformat(), "herkunft": herkunft, "reihen": reihen, "fehler": fehler}
    speicher.schreibe_json(konfig.GESCHICHTE / "preise.json", daten, kompakt=True)
    for kennung, reihe in reihen.items():
        melde(f"Preise {kennung}: {reihe[0]['datum'][:7]} bis {reihe[-1]['datum'][:7]} ({herkunft[kennung]})")
    for f in fehler:
        melde(f"Hinweis Preise: {f}")
    return daten


# ------------------------------------------------------------------ Ernte
def lies_psd(inhalt: bytes, ware: str, laender: list[str]) -> dict[str, dict[str, float]]:
    """USDA PSD (zip mit CSV): Produktion je Land und Wirtschaftsjahr, in 1000 Tonnen."""
    with zipfile.ZipFile(io.BytesIO(inhalt)) as archiv:
        name = next(n for n in archiv.namelist() if n.lower().endswith(".csv"))
        text = archiv.read(name).decode("utf-8-sig", errors="replace")
    ergebnis = {land: {} for land in laender}
    for zeile in csv.DictReader(io.StringIO(text)):
        if (zeile.get("Commodity_Description") == ware and zeile.get("Attribute_Description") == "Production"
                and zeile.get("Country_Name") in ergebnis):
            try:
                ergebnis[zeile["Country_Name"]][zeile["Market_Year"]] = float(zeile["Value"])
            except (KeyError, ValueError):
                continue
    return ergebnis


def ernte(k: dict, melde=print) -> dict:
    g = k["quellen"]["geschichte"]["ernte"]
    reihen = lies_psd(netz.hole_bytes(g["url"]), g["ware"], g["laender"])
    daten = {"stand": date.today().isoformat(), "quelle": g["name"], "einheit": "1000 t", "reihen": reihen}
    speicher.schreibe_json(konfig.GESCHICHTE / "ernte.json", daten, kompakt=True)
    for land, werte in reihen.items():
        jahre = sorted(werte)
        melde(f"Ernte {land}: {jahre[0] if jahre else '–'} bis {jahre[-1] if jahre else '–'} ({len(jahre)} Jahre)")
    return daten


# ------------------------------------------------------------------ Regen
def lies_gz(pfad) -> dict | None:
    if not pfad.exists():
        return None
    try:
        with gzip.open(pfad, "rt", encoding="utf-8") as f:
            return json.load(f)
    except (OSError, json.JSONDecodeError):
        return None


def schreibe_gz(pfad, daten: dict) -> None:
    pfad.parent.mkdir(parents=True, exist_ok=True)
    hilfe = pfad.with_name(pfad.name + ".neu")
    with gzip.open(hilfe, "wt", encoding="utf-8") as f:
        json.dump(daten, f, ensure_ascii=False, separators=(",", ":"), sort_keys=True)
    os.replace(hilfe, pfad)


def regen_pfad(name: str):
    return konfig.GESCHICHTE / "regen" / f"{_slug(name)}.json.gz"


def _abschnitte_ost(ab: date, ende: date) -> list[tuple[date, date, str]]:
    """Fünfjahresblöcke; der erste beginnt bei `ab`, der letzte endet bei `ende`."""
    abschnitte = []
    jahr = ab.year
    while True:
        bis_jahr = (jahr // 5) * 5 + 4
        von = max(ab, date(jahr, 1, 1))
        bis = min(date(bis_jahr, 12, 31), ende)
        abschnitte.append((von, bis, f"{von.year}-{bis_jahr}"))
        if bis >= ende:
            return abschnitte
        jahr = bis_jahr + 1


def _abschnitte_west(ab: date, ende: date) -> list[tuple[date, date, str]]:
    """Je Jahr die kurze Regenzeit Oktober bis Dezember."""
    abschnitte = []
    for jahr in range(ab.year, ende.year + 1):
        von, bis = date(jahr, 10, 1), min(date(jahr, 12, 31), ende)
        if von <= ende and von >= ab:
            abschnitte.append((von, bis, str(jahr)))
    return abschnitte


def regen(k: dict, budget: Budget, heute: date, melde=print) -> dict:
    om = k["quellen"]["openmeteo"]
    ab = _datum(k["quellen"]["geschichte"]["regen_ab"])
    ende = heute - timedelta(days=6)
    stand = {}
    for region in ("ost", "west"):
        for punkt in k["regionen"][region]["punkte"]:
            pfad = regen_pfad(punkt["name"])
            ort = {"lat": punkt["lat"], "lon": punkt["lon"]}
            daten = lies_gz(pfad)
            if not daten or daten.get("ort") != ort:
                daten = {"punkt": punkt["name"], "region": region, "ort": ort, "werte": {}, "fertig": []}
            abschnitte = _abschnitte_ost(ab, ende) if region == "ost" else _abschnitte_west(ab, ende)
            for von, bis, marke in abschnitte:
                if marke in daten["fertig"]:
                    continue
                gewicht = openmeteo.gewicht((bis - von).days + 1)
                budget.pruefe(gewicht, f"Regen {punkt['name']} {marke}")
                try:
                    werte = openmeteo.hole_tage(om["url"], punkt, von, bis, modell=om.get("modell"))
                except AbrufFehler as f:
                    schreibe_gz(pfad, daten)
                    if "429" in str(f):
                        raise Abbruch(f"Open-Meteo-Limit erreicht bei Regen {punkt['name']} {marke}. Morgen weiter.")
                    raise
                finally:
                    budget.verbuche(gewicht)
                daten["werte"].update({d: w for d, w in werte.items() if w is not None})
                if bis <= ende - timedelta(days=60):          # jüngste Wochen bleiben offen für Nachträge
                    daten["fertig"].append(marke)
                schreibe_gz(pfad, daten)
                melde(f"Regen {punkt['name']}: {marke} geholt")
            tage = sorted(daten["werte"])
            stand[punkt["name"]] = {"region": region, "von": tage[0] if tage else None,
                                    "bis": tage[-1] if tage else None, "tage": len(tage)}
    return stand


# ------------------------------------------------------------------ Wind
def wind(k: dict, budget: Budget, heute: date, melde=print) -> dict:
    om = k["quellen"]["openmeteo"]
    monate = k["schwellen"]["dipol"]["wind_monate"]
    ab_jahr = int(k["quellen"]["geschichte"].get("wind_ab", 1960))
    bis_jahr = heute.year if heute >= date(heute.year, monate[-1], 1) + timedelta(days=40) else heute.year - 1
    pfad = konfig.GESCHICHTE / "wind.json"
    daten = speicher.lies_json(pfad, {}) or {}
    daten.setdefault("punkte", {})
    daten.setdefault("orte", {})
    klima = (speicher.lies_json(konfig.KLIMA / "wind.json", {}) or {}).get("punkte", {})
    tage = (date(2001, monate[-1], calendar.monthrange(2001, monate[-1])[1]) - date(2001, monate[0], 1)).days + 1
    for punkt in k["regionen"]["wind_sumatra"]["punkte"]:
        ort = {"lat": punkt["lat"], "lon": punkt["lon"]}
        if daten["orte"].get(punkt["name"], ort) != ort:
            daten["punkte"].pop(punkt["name"], None)            # Punkt verschoben: neu holen
        daten["orte"][punkt["name"]] = ort
        eintrag = daten["punkte"].setdefault(punkt["name"], {"je_jahr": {}})
        schon = set(eintrag["je_jahr"]) | set((klima.get(punkt["name"]) or {}).get("je_jahr", {}))
        for jahr in range(ab_jahr, bis_jahr + 1):
            if str(jahr) in schon:
                continue
            gewicht = openmeteo.gewicht(tage, 2)
            budget.pruefe(gewicht, f"Wind {punkt['name']} {jahr}")
            start = date(jahr, monate[0], 1)
            ende = date(jahr, monate[-1], calendar.monthrange(jahr, monate[-1])[1])
            try:
                mittel, _ = openmeteo.hole_mittleren_zonalwind(om["url"], punkt, start, ende, modell=om.get("modell"))
            except AbrufFehler as f:
                speicher.schreibe_json(pfad, daten)
                if "429" in str(f):
                    raise Abbruch(f"Open-Meteo-Limit erreicht bei Wind {punkt['name']} {jahr}. Morgen weiter.")
                raise
            finally:
                budget.verbuche(gewicht)
            eintrag["je_jahr"][str(jahr)] = round(mittel, 3)
            speicher.schreibe_json(pfad, daten)
        melde(f"Wind {punkt['name']}: {len(eintrag['je_jahr'])} Jahre in der Geschichte, "
              f"{len((klima.get(punkt['name']) or {}).get('je_jahr', {}))} aus der Klimatologie")
    return daten


# ------------------------------------------------------------------ Lauf
def main(argv: list[str] | None = None) -> int:
    teiler = argparse.ArgumentParser(description="Geschichtsdaten für den Rückblick holen (fortsetzbar)")
    teiler.add_argument("--nur", help="Teile mit Komma getrennt: preise, ernte, regen, wind")
    teiler.add_argument("--nur-offene", action="store_true", help="nur Teile, die noch nicht fertig sind")
    teiler.add_argument("--budget-gewicht", type=float, default=9000)
    teiler.add_argument("--budget-minuten", type=float, default=150)
    teiler.add_argument("--pro-minute", type=float, default=None)
    teiler.add_argument("--heute", help="Stichtag JJJJ-MM-TT (nur Tests)")
    args = teiler.parse_args(argv)
    k = konfig.konfiguration()
    heute = date.fromisoformat(args.heute) if args.heute else konfig.jetzt().date()
    pro_minute = args.pro_minute if args.pro_minute is not None else k["quellen"]["openmeteo"]["gewicht_pro_minute"]
    budget = Budget(args.budget_gewicht, args.budget_minuten, pro_minute)
    meldungen: list[str] = []

    def melde(text: str) -> None:
        print(text, flush=True)
        meldungen.append(text)

    statuspfad = konfig.GESCHICHTE / "status.json"
    status = speicher.lies_json(statuspfad, {}) or {}
    alle = ["preise", "ernte", "wind", "regen"]
    teile = [t.strip() for t in args.nur.split(",") if t.strip()] if args.nur else alle
    unbekannt = [t for t in teile if t not in alle]
    if unbekannt:
        teiler.error(f"unbekannte Teile: {', '.join(unbekannt)} (erlaubt: {', '.join(alle)})")
    if args.nur_offene:
        teile = [t for t in teile if not (status.get(t) or {}).get("fertig")]
        if not teile:
            print("Geschichte ist vollständig, nichts zu tun.")
            return 0
    for teil in teile:
        try:
            if teil == "preise":
                d = preise(k, melde)
                status["preise"] = {"stand": d["stand"], "herkunft": d["herkunft"], "fehler": d["fehler"],
                                    "fertig": True}
            elif teil == "ernte":
                d = ernte(k, melde)
                status["ernte"] = {"stand": d["stand"], "fertig": any(d["reihen"].values())}
            elif teil == "wind":
                wind(k, budget, heute, melde)
                status["wind"] = {"stand": heute.isoformat(), "fertig": True}
            else:
                punkte = regen(k, budget, heute, melde)
                status["regen"] = {"stand": heute.isoformat(), "fertig": True, "punkte": punkte}
        except Abbruch as grund:
            status.setdefault(teil, {})["fertig"] = False
            melde(f"Angehalten ({teil}): {grund}")
            break
        except Exception as fehler:        # ein Teil darf die anderen nicht verhindern
            status.setdefault(teil, {})["fertig"] = False
            status[teil]["fehler_heute"] = f"{type(fehler).__name__}: {fehler}"
            melde(f"Fehler bei {teil}: {type(fehler).__name__}: {fehler}")
    status["stand"] = heute.isoformat()
    status["gewicht_heute"] = round(budget.verbraucht)
    melde(f"Verbraucht: etwa {round(budget.verbraucht)} gewichtete Open-Meteo-Abrufe.")
    status["meldungen"] = meldungen                   # lesbar ohne Anmeldung, über die Rohdatei im Repository
    speicher.schreibe_json(statuspfad, status)
    if os.environ.get("GITHUB_STEP_SUMMARY"):
        with open(os.environ["GITHUB_STEP_SUMMARY"], "a", encoding="utf-8") as f:
            f.write("## Geschichtsdaten\n\n" + "\n".join(f"- {m}" for m in meldungen) + "\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
